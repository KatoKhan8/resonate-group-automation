"""TASK-342: Restore the ten's information content in the fifty's workbook.

Parses the posted HTML (503-FIFTY-v2-2026-09-25.html) which carries all the
data, and produces a NEW xlsx with:
  - Sheet 1 "Summary": same metrics as the original
  - Sheet 2 "Leads": expanded with facts-with-sources, hypothesis, signal
    strength, why-that-capability, copylint (which rule fired), sequence gate
    (which rule fired)
  - Sheet 3 "Every message": channel, step, day, thread, new-thread-or-reply,
    subject, objective, body as sent, P.S., signature block, chars

Does NOT overwrite the posted, hashed files.
"""
import re, os, sys, html as htmlmod

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, ROOT)

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill

_CLAUDE_WORKTREE = os.path.join(
    os.path.dirname(ROOT), "resonate-group-automation")
_HTML_CANDIDATES = [
    os.path.join(ROOT, "work", "review", "fifty-verified",
                 "503-FIFTY-v2-2026-09-25.html"),
    os.path.join(_CLAUDE_WORKTREE, "work", "review", "fifty-verified",
                 "503-FIFTY-v2-2026-09-25.html"),
]
HTML_PATH = os.environ.get("FIFTY_HTML") or next(
    (p for p in _HTML_CANDIDATES if os.path.exists(p)), _HTML_CANDIDATES[0])
OUT_PATH = os.environ.get(
    "FIFTY_XLSX_OUT",
    os.path.join(ROOT, "work", "review",
                 "503-FIFTY-v2-2026-09-25-TASK342.xlsx"))

STEPS = ["em1", "em2", "em3", "em4", "em5"]
WAITS = {"em1": 1, "em2": 4, "em3": 8, "em4": 12, "em5": 21}
THREAD = {"em1": "A", "em2": "A", "em3": "B", "em4": "B", "em5": "C"}
LI_STEPS = [("connect", 1, "connection request"),
            ("msg1", 3, "after acceptance"),
            ("msg2", 8, "capability + email cross-reference"),
            ("msg3", 14, "close")]

PS_LABEL = re.compile(r"^\s*p\.?\s*s\.?\s*[:\-]?\s*", re.I)


def unescape(s):
    s = re.sub(r"<br\s*/?>", "\n", s)
    s = re.sub(r"&middot;", "·", s)
    s = re.sub(r"&ldquo;", "\u201c", s)
    s = re.sub(r"&rdquo;", "\u201d", s)
    s = re.sub(r"&amp;", "&", s)
    s = re.sub(r"&#x27;", "'", s)
    s = re.sub(r"&lt;", "<", s)
    s = re.sub(r"&gt;", ">", s)
    s = re.sub(r"&uarr;", "↑", s)
    s = re.sub(r"<[^>]+>", "", s)
    return s.strip()


def strip_tags(s):
    return re.sub(r"<[^>]+>", "", s).strip()


def parse_lead(section):
    rec = {}
    header_m = re.search(
        r"<header><h1>(.+?)</h1><p>(.+?)</p></header>", section, re.S)
    if not header_m:
        return None
    rec["lead_name"] = unescape(header_m.group(1))
    meta = unescape(header_m.group(2))
    meta_parts = [p.strip() for p in meta.split("·")]
    rec["email"] = meta_parts[0].strip() if len(meta_parts) > 0 else ""
    rec["title"] = meta_parts[1].strip() if len(meta_parts) > 1 else ""
    sender_m = re.search(r"from\s+(.+?)\s+<(.+?)>", meta)
    if sender_m:
        rec["sender_name"] = sender_m.group(1).strip()
        rec["sender_email"] = sender_m.group(2).strip()
    else:
        rec["sender_name"] = ""
        rec["sender_email"] = ""
    cohort_m = re.search(r"(US-\w+ / [^·↑]+)", meta)
    rec["cohort"] = cohort_m.group(1).strip() if cohort_m else ""

    company_m = re.search(r"—\s*(.+?)$",
                          re.sub(r"<[^>]+>", "", header_m.group(1)))
    rec["company"] = company_m.group(1).strip() if company_m else ""

    li_m = re.search(r"LinkedIn:\s*(https?://\S+)", section[:3000])
    rec["linkedin"] = li_m.group(1) if li_m else "(none)"
    if "no LinkedIn URL" in section[:3000]:
        rec["linkedin"] = "(none)"

    qual_m = re.search(
        r"class='tag\s*([^']*)'>(QUALIFIED_\w+|UNQUALIFIED|INSUFFICIENT|HELD)",
        section)
    rec["qualification"] = qual_m.group(2) if qual_m else "?"

    ps_m = re.search(r"class='tag v'>P\.S\.\s*([^<]+)</span>", section)
    rec["ps_variant"] = ps_m.group(1).strip() if ps_m else "-"

    icp_m = re.search(
        r"<b>ICP \(stage A\):</b>\s*(.+?)<div class='meta'[^>]*>(.+?)</div>",
        section, re.S)
    if icp_m:
        rec["icp_verdict"] = unescape(icp_m.group(1))
        rec["icp_evidence"] = unescape(icp_m.group(2))
    else:
        rec["icp_verdict"] = "-"
        rec["icp_evidence"] = "-"

    hyp_m = re.search(
        r"class='hyplabel'>HYPOTHESIS.*?</div><b>(.+?)</b>"
        r".*?rests on:\s*(.+?)\s*·\s*signal:\s*<b>(.+?)</b>",
        section, re.S)
    if hyp_m:
        rec["hypothesis"] = unescape(hyp_m.group(1))
        rec["hypothesis_rests_on"] = unescape(hyp_m.group(2))
        rec["signal_strength"] = unescape(hyp_m.group(3))
    else:
        rec["hypothesis"] = "-"
        rec["hypothesis_rests_on"] = "-"
        rec["signal_strength"] = "-"

    cap_m = re.search(
        r"<b>capability chosen:</b>\s*<span class='tag v'>([^<]+)</span>"
        r".*?<b>why_this_one|why.*?:</b>\s*(.+?)<br>"
        r"<b>what changes for them:</b>\s*(.+?)</div>",
        section, re.S)
    if not cap_m:
        cap_m = re.search(
            r"<b>capability chosen:</b>.*?tag v'>([^<]+)</span>"
            r"<div class='meta'[^>]*>(.+?)<br><b>what changes for them:</b>\s*"
            r"(.+?)</div>",
            section, re.S)
    if cap_m:
        rec["capability"] = unescape(cap_m.group(1))
        rec["why_capability"] = unescape(cap_m.group(2))
        rec["what_changes"] = unescape(cap_m.group(3))
    else:
        rec["capability"] = "-"
        rec["why_capability"] = "-"
        rec["what_changes"] = "-"

    held_m = re.search(
        r"class='held'><b>HELD.*?</b>.*?<div class='meta'[^>]*>(.+?)</div>",
        section, re.S)
    rec["held_reason"] = unescape(held_m.group(1)) if held_m else ""

    seq_gate_m = re.search(
        r"class='gate[^']*'><b>sequence gate:\s*(.+?)</b>"
        r"(?:<pre>(.+?)</pre>)?",
        section, re.S)
    if seq_gate_m:
        rec["seq_gate_verdict"] = unescape(seq_gate_m.group(1))
        rec["seq_gate_detail"] = unescape(seq_gate_m.group(2) or "")
    else:
        rec["seq_gate_verdict"] = "-"
        rec["seq_gate_detail"] = ""

    cl_m = re.search(
        r"class='gate[^']*'><b>copylint:\s*(.+?)</b>"
        r"(?:<div class='meta'>(.+?)</div>)?",
        section, re.S)
    if cl_m:
        rec["copylint_verdict"] = unescape(cl_m.group(1))
        fired = unescape(cl_m.group(2) or "")
        rec["copylint_fired"] = fired if fired != "nothing fired" else "(clean)"
    else:
        rec["copylint_verdict"] = "-"
        rec["copylint_fired"] = "-"

    facts = []
    for fm in re.finditer(
            r"class='fact'><b>(.+?)</b><div class='u'>(.+?)</div>"
            r"<div class='q'>(.+?)</div>",
            section, re.S):
        facts.append({
            "text": unescape(fm.group(1)),
            "source": unescape(fm.group(2)),
            "quote": unescape(fm.group(3)),
        })
    rec["facts"] = facts

    emails = []
    for em in re.finditer(
            r"class='msg'>.*?class='day'>day (\d+)</span>"
            r"<span>(.+?)</span>.*?"
            r"(?:class='obj'><b>objective:</b>\s*(.+?)</div>)?"
            r"(?:<div class='subj'>(.+?)</div>)?"
            r"<div class='bd'>(.+?)</div>",
            section, re.S):
        day = int(em.group(1))
        thread_info = unescape(em.group(2))
        objective = unescape(em.group(3) or "")
        subject_line = unescape(em.group(4) or "")
        body = unescape(em.group(5))
        step_key = None
        for k, d in WAITS.items():
            if d == day:
                step_key = k
                break
        ps_m2 = re.search(
            r"class='ps'>P\.S\.\s*(.+?)</div>",
            section[section.find(em.group(0)):section.find(em.group(0)) + 3000],
            re.S)
        ps_text = unescape(ps_m2.group(1)) if ps_m2 else ""
        sig_m2 = re.search(
            r"class='siglabel'>appended by the mailbox</div>(.+?)</div>",
            section[section.find(em.group(0)):section.find(em.group(0)) + 5000],
            re.S)
        sig_text = unescape(sig_m2.group(1)) if sig_m2 else ""
        is_reply = "reply" in thread_info.lower()
        thread_id = ""
        tid_m = re.search(r"thread\s+([A-C])", thread_info, re.I)
        if tid_m:
            thread_id = tid_m.group(1)
        emails.append({
            "channel": "email",
            "step": step_key or ("day%d" % day),
            "day": day,
            "thread": thread_id,
            "new_or_reply": "reply in thread %s" % thread_id if is_reply
                            else "NEW THREAD %s" % thread_id,
            "subject": subject_line,
            "objective": objective,
            "body": body,
            "ps": ps_text,
            "signature": sig_text,
            "chars": len(body),
        })

    li_msgs = []
    li_section = section[section.find("LinkedIn cadence"):]
    if li_section:
        for lm in re.finditer(
                r"class='msg'>.*?class='day'>day (\d+)</span>"
                r"<span>(.+?)\s*·\s*(\d+)\s*chars",
                li_section, re.S):
            day = int(lm.group(1))
            label = unescape(lm.group(2))
            chars = int(lm.group(3))
            obj_m = re.search(
                r"class='obj'><b>objective:</b>\s*(.+?)</div>",
                li_section[li_section.find(lm.group(0)):
                           li_section.find(lm.group(0)) + 1000], re.S)
            obj = unescape(obj_m.group(1)) if obj_m else ""
            bd_m = re.search(
                r"<div class='bd'>(.+?)</div>",
                li_section[li_section.find(lm.group(0)):
                           li_section.find(lm.group(0)) + 3000], re.S)
            body = unescape(bd_m.group(1)) if bd_m else ""
            step_key = ""
            for k, d, _ in LI_STEPS:
                if d == day:
                    step_key = k
                    break
            li_msgs.append({
                "channel": "linkedin",
                "step": step_key or ("day%d" % day),
                "day": day,
                "thread": "",
                "new_or_reply": label,
                "subject": "",
                "objective": obj,
                "body": body,
                "ps": "",
                "signature": "",
                "chars": len(body) if body else chars,
            })

    rec["messages"] = emails + li_msgs
    return rec


def facts_text(facts):
    if not facts:
        return "(no facts extracted)"
    out = []
    for i, f in enumerate(facts, 1):
        out.append('%d. %s\n   source: %s\n   verbatim: "%s"'
                   % (i, f["text"], f["source"], f["quote"][:220]))
    return "\n".join(out)


def main():
    raw = open(HTML_PATH, encoding="utf-8").read()
    parts = re.split(
        r"<div style='border-top:6px solid #1f3864;margin-top:30px'></div>",
        raw)
    leads = []
    for p in parts[1:]:
        rec = parse_lead(p)
        if rec:
            leads.append(rec)
    print("parsed %d leads from HTML" % len(leads))

    written = sum(1 for r in leads if r["messages"])
    held = sum(1 for r in leads if r["held_reason"])
    print("written %d  held %d" % (written, held))

    wb = Workbook()

    # --- Sheet 1: Summary ---
    ws1 = wb.active
    ws1.title = "Summary"
    ws1.append(["Metric", "Value"])
    for c in ws1[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3864")
    caps = set()
    for r in leads:
        if r["capability"] != "-":
            caps.add(r["capability"])
    quals = {}
    for r in leads:
        q = r["qualification"]
        quals[q] = quals.get(q, 0) + 1
    cl_clean = sum(1 for r in leads if r["copylint_verdict"] == "clean")
    cl_refused = sum(1 for r in leads
                     if r["copylint_verdict"] in ("refused", "REFUSED"))
    sg_pass = sum(1 for r in leads if "PASSED" in r["seq_gate_verdict"])
    sg_fail = sum(1 for r in leads if "FAILED" in r["seq_gate_verdict"])
    both = sum(1 for r in leads
               if r["copylint_verdict"] == "clean"
               and "PASSED" in r["seq_gate_verdict"])

    ws1.append(["Total leads", len(leads)])
    ws1.append(["Written", written])
    ws1.append(["Held", held])
    ws1.append(["Distinct capabilities", len(caps)])
    ws1.append(["Capabilities used", ", ".join(sorted(caps))])
    ws1.append([])
    ws1.append(["Gate breakdown", ""])
    ws1.append(["copylint clean", cl_clean])
    ws1.append(["copylint refused", cl_refused])
    ws1.append(["sequence gate passed", sg_pass])
    ws1.append(["sequence gate failed", sg_fail])
    ws1.append(["passing BOTH", both])
    ws1.append([])
    ws1.append(["Qualification", "Count"])
    for q in sorted(quals):
        ws1.append([q, quals[q]])
    for row in ws1.iter_rows(min_row=2, max_row=ws1.max_row):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)
    ws1.column_dimensions["A"].width = 28
    ws1.column_dimensions["B"].width = 50

    # --- Sheet 2: Leads (expanded) ---
    ws2 = wb.create_sheet("Leads")
    HEAD2 = [
        "company", "lead", "email", "title", "sender mailbox", "sender name",
        "LinkedIn URL", "qualification", "cohort",
        "ICP verdict", "ICP evidence",
        "facts with sources",
        "HYPOTHESIS (not a fact)", "hypothesis rests on", "signal strength",
        "capability chosen", "why that capability", "what changes for them",
        "P.S. variant",
        "copylint (verdict)", "copylint (rule fired)",
        "sequence gate (verdict)", "sequence gate (detail)",
        "held / error reason",
    ]
    ws2.append(HEAD2)
    for c in ws2[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3864")
        c.alignment = Alignment(vertical="top", wrap_text=True)

    for rec in leads:
        ws2.append([
            rec["company"],
            rec["lead_name"],
            rec["email"],
            rec["title"],
            rec["sender_email"],
            rec["sender_name"],
            rec["linkedin"],
            rec["qualification"],
            rec["cohort"],
            rec["icp_verdict"],
            rec["icp_evidence"],
            facts_text(rec["facts"]),
            rec["hypothesis"],
            rec["hypothesis_rests_on"],
            rec["signal_strength"],
            rec["capability"],
            rec["why_capability"],
            rec["what_changes"],
            rec["ps_variant"],
            rec["copylint_verdict"],
            rec["copylint_fired"],
            rec["seq_gate_verdict"],
            rec["seq_gate_detail"],
            rec["held_reason"],
        ])
        for c in ws2[ws2.max_row]:
            c.alignment = Alignment(vertical="top", wrap_text=True)

    for col, w in {"A": 22, "B": 20, "C": 28, "D": 24, "E": 30, "F": 18,
                   "G": 38, "H": 17, "I": 28}.items():
        ws2.column_dimensions[col].width = w
    for col in "JKLMNOPQRSTUVWX":
        ws2.column_dimensions[col].width = 52
    ws2.column_dimensions["L"].width = 70
    ws2.freeze_panes = "B2"

    # --- Sheet 3: Every message ---
    ws3 = wb.create_sheet("Every message")
    HEAD3 = ["company", "channel", "step", "day", "thread",
             "new thread or reply", "subject", "objective", "body (as sent)",
             "P.S.", "signature block (appended by the mailbox)", "chars"]
    ws3.append(HEAD3)
    for c in ws3[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3864")
        c.alignment = Alignment(vertical="top", wrap_text=True)

    for rec in leads:
        if not rec["messages"]:
            ws3.append([rec["company"], "-", "HELD", "", "", "", "", "",
                        rec["held_reason"], "", "", 0])
            continue
        for msg in rec["messages"]:
            ws3.append([
                rec["company"],
                msg["channel"],
                msg["step"],
                msg["day"],
                msg["thread"],
                msg["new_or_reply"],
                msg["subject"],
                msg["objective"],
                msg["body"],
                msg["ps"],
                msg["signature"],
                msg["chars"],
            ])
            for c in ws3[ws3.max_row]:
                c.alignment = Alignment(vertical="top", wrap_text=True)

    for col, w in {"A": 20, "B": 10, "C": 9, "D": 6, "E": 8, "F": 22,
                   "G": 34, "H": 44, "I": 90, "J": 44, "K": 34, "L": 7}.items():
        ws3.column_dimensions[col].width = w
    ws3.freeze_panes = "A2"

    os.makedirs(os.path.dirname(OUT_PATH) or ".", exist_ok=True)
    wb.save(OUT_PATH)
    print(OUT_PATH, os.path.getsize(OUT_PATH), "bytes")
    print("sheet 1 (Summary): %d rows" % (ws1.max_row - 1))
    print("sheet 2 (Leads): %d rows" % (ws2.max_row - 1))
    print("sheet 3 (Every message): %d rows" % (ws3.max_row - 1))

    max_cell = 0
    for row in ws3.iter_rows(min_row=2, max_row=ws3.max_row):
        for c in row:
            if c.value and len(str(c.value)) > max_cell:
                max_cell = len(str(c.value))
    print("longest cell in messages sheet: %d chars" % max_cell)


if __name__ == "__main__":
    main()
