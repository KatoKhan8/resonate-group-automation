#!/usr/bin/env python3
"""GATE 1. The file a person reads BEFORE a campaign is ever activated.

    py -m src.reviewfile 503 --snapshot work/review/raw/bison-503.json

Writes `work/review/<campaign>-<date>.xlsx` and `work/review/<campaign>-
<date>.html`: ONE ROW PER LEAD, carrying the sender mailbox, the sender's
NAME, and the subject and full body of every step EXACTLY AS THE PROVIDER
WILL SEND IT - plus the pack fact the opener quotes and the URL it came
from, and the verdict of gates 2 and 3 on every step.

Operator's standing directive, 2026-09-25.

## IT READS FROM THE PROVIDER. THAT IS THE ENTIRE POINT.

A review file built from our own render proves that our render is what our
render is. The push that caused the incident never rendered anything: it
POSTed hand-written string literals straight at `bison.create_lead`, so
there was no render to check and every local gate would have passed a file
generated from one. What went to sixty-four people was what the PROVIDER
held, and that is the only thing worth printing.

So every cell comes from a provider read, and each row says which one:

  `provider_queue`      `GET /campaigns/{id}/scheduled-emails` - the
                        pre-send queue, merge fields already resolved. This
                        is the strongest evidence there is: it is the row
                        the provider will hand to the mailbox.
  `provider_rendered`   the campaign's stored sequence template (`GET
                        /campaigns/{id}/sequence-steps`) filled from the
                        LEAD'S stored custom variables (`GET
                        /campaigns/{id}/leads`). Used where the provider has
                        not built a queue row yet - a stopped campaign has
                        no queue for its later steps - and still two
                        provider reads rather than one of ours.

A cell is NEVER left blank as though it were fine. A step with no copy at
the provider prints `** NO COPY AT THE PROVIDER **`, which is what 46 of
campaign 491's 333 leads would print today.

## WHERE IT MAY BE WRITTEN

Under `work/`, which is gitignored, and nowhere else. These files carry
real recipients - 64 of them are people who have already received the wrong
email - and `refuse_outside_work` raises rather than trusting this
paragraph to be read. `docs/` is named explicitly in the refusal because
that is the directory somebody would reach for.
"""
import argparse
import datetime
import html
import json
import os
import re
import sys

from . import copyprovenance, export, packfact

#: Where a review file may be written. Relative to the repository root.
REVIEW_DIR = os.path.join("work", "review")

#: What a cell says when the provider holds nothing. Never blank: an empty
#: cell reads as "fine" and this reads as what it is.
NO_COPY = "** NO COPY AT THE PROVIDER **"
NOT_RECORDED = "** NOT RECORDED AT THE PROVIDER **"

#: How the opener names the line it is quoting. Both shapes that have
#: actually shipped. Used ONLY to recover the span from copy that is
#: already at the provider - a lead staged through `copyprovenance.certify`
#: carries the span, the fact id and the source url as custom variables and
#: nothing has to be recovered.
QUOTED_SPAN_PATTERNS = (
    re.compile(r"the line about (.+?) is what made me write", re.S | re.I),
    re.compile(r"your site says (.+?)[.\n]", re.S | re.I),
    re.compile(r"I was reading .{0,60}? and (?:the line about )?(.+?) is what",
               re.S | re.I),
)

#: The custom variables a certified lead carries about its pack fact.
SPAN_VARIABLE = "pack_fact_span"
FACT_ID_VARIABLE = "pack_fact_id"
SOURCE_URL_VARIABLE = "pack_fact_source_url"

_MERGE = re.compile(r"\{([A-Za-z_][A-Za-z0-9_]*)\}")


class ReviewRefused(RuntimeError):
    """A review file was about to be written where it must not be."""


def root():
    return os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))


def refuse_outside_work(path):
    """Raise unless `path` is inside `work/`. Called before every write.

    THE FILTER, AND IT RUNS BEFORE THE FIRST BYTE. A review file is 64 real
    recipients with their addresses, their companies and the exact words
    that were sent to them. `work/` is gitignored; `docs/` is committed and
    pushed to GitHub. Those two facts are the whole reason this function
    exists, so `docs/` is named in the refusal rather than left to be
    inferred from "outside work/".

    Resolved with `realpath` on both sides, so a symlink or a `..` cannot
    spell a path that passes here and lands elsewhere.
    """
    target = os.path.realpath(os.path.abspath(path))
    allowed = os.path.realpath(os.path.join(root(), "work"))
    if target == allowed or target.startswith(allowed + os.sep):
        return target
    raise ReviewRefused(
        "a review file was about to be written to %s. It carries real "
        "recipients and the exact words sent to them, so it may only be "
        "written under %s, which is gitignored. docs/ is committed and "
        "pushed; a review file there publishes a client's prospect list."
        % (target, allowed))


def out_dir():
    return os.path.join(root(), REVIEW_DIR)


# --------------------------------------------------------------- provider view

def variables_of(lead):
    return copyprovenance.variables_of(lead)


def _render(template, variables, lead):
    """One stored sequence step filled from one stored lead. Provider x provider.

    An unresolved merge field is left VISIBLE as `{BODY_4}` rather than
    blanked, because a blank is how a step with no copy reads as a step with
    short copy.
    """
    facts = {k.upper(): v for k, v in (variables or {}).items()}
    for key in ("first_name", "last_name", "company", "email"):
        if (lead or {}).get(key) is not None:
            facts[key.upper()] = lead.get(key)

    def sub(match):
        name = match.group(1).upper()
        return str(facts[name]) if name in facts else match.group(0)

    return _MERGE.sub(sub, str(template or ""))


def queue_by_lead(snapshot):
    """`{lead_id: {order: queue_row}}` plus the sender each lead is bound to.

    ORDER COMES FROM THE SEQUENCE, NOT FROM THE QUEUE. A queue row names
    `sequence_step_id`, which is a provider id, so it is joined back to the
    stored sequence to learn which step of the cadence it is. Reading the
    queue's own arrival order instead would number a campaign's steps by
    when the provider happened to build them.
    """
    order_of = {}
    for position, step in enumerate(snapshot.get("sequence") or [], start=1):
        if step.get("id") is not None:
            order_of[str(step["id"])] = step.get("order") or position
    out, senders = {}, {}
    for row in snapshot.get("queue") or []:
        lead_id = ((row.get("lead") or {}) or {}).get("id")
        if lead_id is None:
            continue
        position = order_of.get(str(row.get("sequence_step_id")))
        out.setdefault(str(lead_id), {})[position] = row
        sender = row.get("sender_email") or {}
        if sender.get("email"):
            senders.setdefault(str(lead_id), set()).add(
                (sender.get("name") or "", sender.get("email") or ""))
    return out, senders


def sender_pool(snapshot):
    """`[(name, email)]` for every mailbox the campaign is bound to.

    Used only when the provider has not yet built a queue row for a lead, so
    which of them will send is genuinely not decided. The review file says
    exactly that instead of naming one - and `check_step` is handed NO owner
    for such a lead, so any signature on it is refused for being
    uncomparable. That is the fail-closed direction: an unbound lead is the
    case where a constant signature is least visible.
    """
    out = []
    for row in snapshot.get("sender_pool") or []:
        if isinstance(row, dict) and row.get("email"):
            out.append((row.get("name") or "", row.get("email")))
    return sorted(set(out))


def steps_of(snapshot, lead, queue):
    """Every step for one lead, as the provider will send it.

    `[{order, subject, body, status, source}]`. The queue row wins whenever
    there is one; otherwise the stored sequence template is filled from the
    lead's stored variables.
    """
    variables = variables_of(lead)
    sequence = snapshot.get("sequence") or []
    out = []
    for position, step in enumerate(sequence, start=1):
        order = step.get("order") or position
        row = (queue or {}).get(order)
        if row is not None:
            out.append({"order": order,
                        "subject": row.get("email_subject") or "",
                        "body": row.get("email_body") or "",
                        "status": row.get("status") or "",
                        "scheduled": row.get("scheduled_date") or "",
                        "source": "provider_queue"})
            continue
        out.append({"order": order,
                    "subject": _render(step.get("email_subject"), variables, lead),
                    "body": _render(step.get("email_body"), variables, lead),
                    "status": "",
                    "scheduled": "",
                    "source": "provider_rendered"})
    return out


# ------------------------------------------------------------------ pack fact

def quoted_span(text):
    """The span this opener claims to have read on their site, or None."""
    plain = copyprovenance.plain(text)
    for pattern in QUOTED_SPAN_PATTERNS:
        match = pattern.search(plain)
        if match:
            return re.sub(r"\s+", " ", match.group(1)).strip(" .,:;")
    return None


def fact_for(span, pack):
    """The pack fact a span was read off, or None.

    Matched by CONTAINMENT in the fact's own snippet, which is the only
    honest join available for copy that is already at the provider: the
    lead carries no fact id, so the question "where did this come from" can
    only be answered by finding the page whose text contains it. A span
    found in no fact is reported as exactly that, never as unattributed.
    """
    needle = re.sub(r"\s+", " ", str(span or "")).strip().lower()
    if not needle:
        return None
    for fact in packfact.facts_of(pack):
        snippet = re.sub(r"\s+", " ", str(fact.get("snippet") or "")).lower()
        if needle and needle in snippet:
            return fact
    return None


def pack_for(lead, packs):
    """This lead's research pack, by record id or by email domain."""
    if not packs:
        return None
    variables = variables_of(lead)
    for key in (variables.get("record_id"), (lead or {}).get("company")):
        if key and str(key) in packs:
            return packs[str(key)]
    email = str((lead or {}).get("email") or "")
    if "@" in email:
        domain = email.split("@", 1)[1].lower()
        if domain in packs:
            return packs[domain]
    record = str(variables.get("record_id") or "")
    # `record_id` is the domain with dots replaced by hyphens on some
    # cohorts (`jairusmarketing-com`) and the bare domain on others.
    guess = record.replace("-com", ".com").replace("-", ".")
    return packs.get(record) or packs.get(guess)


# ---------------------------------------------------------------------- rows

def rows(snapshot, *, packs=None, config=None, client=None,
         expect_signature=None):
    """One row per lead, with every step and every verdict."""
    approved = copyprovenance.template_ids(config, client)
    queues, senders = queue_by_lead(snapshot)
    pool = sender_pool(snapshot)
    unbound_name = ("** NOT BOUND YET: one of %d mailboxes (%s) **"
                    % (len(pool),
                       ", ".join(sorted({n for n, _e in pool if n})[:6]))
                    if pool else NOT_RECORDED)
    out = []
    for lead in snapshot.get("leads") or []:
        lead_id = str(lead.get("id"))
        variables = variables_of(lead)
        bound = sorted(senders.get(lead_id) or ())
        sender_names = [n for n, _e in bound if n]
        sender_name = " / ".join(sorted(set(sender_names))) or unbound_name
        sender_email = " / ".join(e for _n, e in bound) or unbound_name
        steps = steps_of(snapshot, lead, queues.get(lead_id))

        span = None
        for step in steps:
            span = variables.get(SPAN_VARIABLE) or quoted_span(step["body"])
            if span:
                break
        pack = pack_for(lead, packs)
        fact = fact_for(span, pack) if span else None
        if not fact and variables.get(FACT_ID_VARIABLE):
            fact = {"fact_id": variables.get(FACT_ID_VARIABLE),
                    "source_url": variables.get(SOURCE_URL_VARIABLE),
                    "snippet": span}
        gate3 = packfact.check_span(span, fact)

        gate2_reasons, gate3_reasons = [], []
        if not gate3["ok"]:
            gate3_reasons = list(gate3["reasons"])
        step_cells = []
        for step in steps:
            position = step["order"]
            if not copyprovenance.plain(step["body"]).strip():
                gate2_reasons.append("step %s: %s" % (position, NO_COPY))
                step_cells.append(dict(step, subject=step["subject"] or NO_COPY,
                                       body=NO_COPY))
                continue
            verdict = copyprovenance.check_step(
                step["body"],
                template_id=variables.get(
                    copyprovenance.TEMPLATE_VARIABLE % position),
                owner_name=sender_name if bound else "",
                approved_ids=approved, subject=step["subject"],
                expect_signature=expect_signature)
            for why in verdict["reasons"]:
                gate2_reasons.append("step %s: %s" % (position, why))
            step_cells.append(dict(step, signature=verdict["signature"]))

        out.append({
            "lead_id": lead_id,
            "email": lead.get("email") or "",
            "first_name": lead.get("first_name") or "",
            "last_name": lead.get("last_name") or "",
            "company": lead.get("company") or "",
            "record_id": variables.get("record_id") or "",
            "contact_key": variables.get("contact_key") or "",
            "sender_mailbox": sender_email,
            "sender_name": sender_name,
            "pack_fact": span or NOT_RECORDED,
            "pack_fact_source_url": (fact or {}).get("source_url") or NOT_RECORDED,
            "pack_fact_id": (fact or {}).get("fact_id") or NOT_RECORDED,
            "gate2_ok": not gate2_reasons,
            "gate2_reasons": gate2_reasons,
            "gate3_ok": gate3["ok"],
            "gate3_reasons": gate3_reasons,
            "verdict": ("SEND" if not gate2_reasons and gate3["ok"]
                        else "HOLD"),
            "steps": step_cells,
        })
    return out


# -------------------------------------------------------------------- writing

def _columns(rows_):
    steps = max([len(r["steps"]) for r in rows_] or [0])
    header = ["lead_id", "email", "first_name", "last_name", "company",
              "record_id", "contact_key", "sender_mailbox", "sender_name",
              "verdict", "gate2", "gate2_reasons", "gate3", "gate3_reasons",
              "pack_fact", "pack_fact_source_url", "pack_fact_id"]
    for n in range(1, steps + 1):
        header += ["step%d_status" % n, "step%d_source" % n,
                   "step%d_subject" % n, "step%d_body" % n]
    return header, steps


def _cells(row, steps):
    out = [row["lead_id"], row["email"], row["first_name"], row["last_name"],
           row["company"], row["record_id"], row["contact_key"],
           row["sender_mailbox"], row["sender_name"], row["verdict"],
           "PASS" if row["gate2_ok"] else "FAIL", " | ".join(row["gate2_reasons"]),
           "PASS" if row["gate3_ok"] else "FAIL", " | ".join(row["gate3_reasons"]),
           row["pack_fact"], row["pack_fact_source_url"], row["pack_fact_id"]]
    by_order = {s["order"]: s for s in row["steps"]}
    for n in range(1, steps + 1):
        step = by_order.get(n) or {}
        out += [step.get("status", ""), step.get("source", ""),
                step.get("subject", ""),
                copyprovenance.plain(step.get("body", ""))]
    return out


def write_xlsx(path, rows_):
    """The spreadsheet. Every cell through `export.safe_cell`.

    A prospect's company name is attacker-controlled text and a leading `=`
    is a formula in Excel exactly as it is in a CSV - `src/export.py` says
    why at length. The guard is applied here for the same reason and not
    because a CSV happened to be the format it was written for.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    header, steps = _columns(rows_)
    book = Workbook()
    sheet = book.active
    sheet.title = "review"
    sheet.append([export.safe_cell(h) for h in header])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows_:
        sheet.append([export.safe_cell(c) for c in _cells(row, steps)])
    widths = {"email": 34, "company": 26, "sender_mailbox": 34,
              "sender_name": 22, "gate2_reasons": 60, "gate3_reasons": 60,
              "pack_fact": 60, "pack_fact_source_url": 40}
    for index, name in enumerate(header, start=1):
        letter = sheet.cell(row=1, column=index).column_letter
        if name.endswith("_body"):
            sheet.column_dimensions[letter].width = 90
        elif name.endswith("_subject"):
            sheet.column_dimensions[letter].width = 45
        else:
            sheet.column_dimensions[letter].width = widths.get(name, 16)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    sheet.freeze_panes = "A2"
    refuse_outside_work(path)
    book.save(path)
    return path


_HTML_HEAD = """<!doctype html><meta charset="utf-8">
<title>%(title)s</title>
<style>
 body{font:13px/1.5 -apple-system,Segoe UI,Roboto,sans-serif;margin:24px;color:#111}
 h1{font-size:18px} h2{font-size:15px;margin:28px 0 6px}
 .lead{border:1px solid #ddd;border-radius:6px;padding:12px 14px;margin:14px 0}
 .HOLD{border-left:5px solid #c00} .SEND{border-left:5px solid #0a0}
 .meta{color:#555;font-size:12px}
 .why{color:#c00;font-size:12px;margin:4px 0}
 .step{background:#fafafa;border:1px solid #eee;padding:8px 10px;margin:8px 0;
       white-space:pre-wrap}
 .subj{font-weight:600} .src{color:#888;font-size:11px}
 .fact{background:#fffbe6;border:1px solid #f0e0a0;padding:6px 8px;font-size:12px}
</style>
<h1>%(title)s</h1>
<p class=meta>%(summary)s</p>
<p class=meta><b>Read back from the provider</b>, not from our render.
 <code>provider_queue</code> is the pre-send row the provider will hand to the
 mailbox; <code>provider_rendered</code> is the campaign's stored sequence
 filled from the lead's stored custom variables.</p>
<p class=meta>This file carries real recipients. It lives under
 <code>work/</code>, which is gitignored, and must never be copied into
 <code>docs/</code>.</p>
"""


def write_html(path, rows_, title, summary):
    esc = html.escape
    parts = [_HTML_HEAD % {"title": esc(title), "summary": esc(summary)}]
    for row in rows_:
        parts.append('<div class="lead %s">' % row["verdict"])
        parts.append("<h2>%s &mdash; %s &lt;%s&gt;</h2>"
                     % (esc(row["verdict"]),
                        esc(("%s %s" % (row["first_name"], row["last_name"])).strip()
                            or row["email"]),
                        esc(row["email"])))
        parts.append('<p class=meta>company %s &middot; lead %s &middot; '
                     'from <b>%s</b> &lt;%s&gt;</p>'
                     % (esc(row["company"]), esc(row["lead_id"]),
                        esc(row["sender_name"]), esc(row["sender_mailbox"])))
        parts.append('<p class=fact>pack fact: %s<br>source: %s</p>'
                     % (esc(row["pack_fact"]),
                        esc(row["pack_fact_source_url"])))
        for why in row["gate2_reasons"] + row["gate3_reasons"]:
            parts.append('<p class=why>%s</p>' % esc(why))
        for step in row["steps"]:
            parts.append('<div class=step><span class=subj>%s</span>'
                         '<span class=src> [%s %s]</span>\n\n%s</div>'
                         % (esc(str(step.get("subject") or "")),
                            esc(str(step.get("source") or "")),
                            esc(str(step.get("status") or "")),
                            esc(copyprovenance.plain(step.get("body") or ""))))
        parts.append("</div>")
    refuse_outside_work(path)
    with open(path, "w", encoding="utf-8", newline="\n") as f:
        f.write("\n".join(parts))
    return path


def write(campaign, rows_, *, directory=None, date=None):
    """Both files, named `<campaign>-<date>`. Returns their paths."""
    directory = directory or out_dir()
    refuse_outside_work(directory)
    os.makedirs(directory, exist_ok=True)
    date = date or datetime.date.today().isoformat()
    stem = os.path.join(directory, "%s-%s" % (campaign, date))
    holds = sum(1 for r in rows_ if r["verdict"] == "HOLD")
    summary = ("%d leads, %d HOLD, %d SEND. Gate 2 (copy provenance) fails "
               "on %d; gate 3 (pack fact) fails on %d."
               % (len(rows_), holds, len(rows_) - holds,
                  sum(1 for r in rows_ if not r["gate2_ok"]),
                  sum(1 for r in rows_ if not r["gate3_ok"])))
    return {"xlsx": write_xlsx(stem + ".xlsx", rows_),
            "html": write_html(stem + ".html", rows_,
                               "Review - campaign %s - %s" % (campaign, date),
                               summary),
            "summary": summary}


def load_packs(paths):
    """A `{key: pack}` index over research-pack files, keyed every way a
    lead might name itself.

    REFUSES A SUMMARY FILE. `work/researchpack-us-cohort-2026-09-25.jsonl`
    carries `"facts": 3` - an integer count - and sits beside three files
    with nearly the same name that carry the real list. Indexing it would
    put rows into the pack index that answer "yes I have three facts" and
    hold none, which is the absent-field trap with a number in it.
    """
    index = {}
    for path in paths or ():
        if not os.path.exists(path):
            continue
        for line in open(path, encoding="utf-8"):
            if not line.strip():
                continue
            row = json.loads(line)
            if not isinstance(row.get("facts"), list):
                continue
            domain = str(row.get("domain") or "").lower()
            if not domain:
                continue
            index[domain] = row
            index[domain.replace(".", "-")] = row
    return index


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("campaign", help="EmailBison campaign id")
    ap.add_argument("--snapshot", help="a file from scripts/copy_snapshot.py")
    ap.add_argument("--client", default="productive")
    ap.add_argument("--packs", nargs="*", default=())
    ap.add_argument("--date")
    a = ap.parse_args(argv)

    from . import clients
    if a.snapshot:
        with open(a.snapshot, encoding="utf-8") as f:
            snapshot = json.load(f)
    else:
        sys.path.insert(0, root())
        from scripts import copy_snapshot
        snapshot = copy_snapshot.snapshot(a.campaign)
    config = clients.load(a.client)
    result = write(a.campaign,
                   rows(snapshot, packs=load_packs(a.packs), config=config,
                        client=a.client),
                   date=a.date)
    print(result["summary"])
    print(result["xlsx"])
    print(result["html"])
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
